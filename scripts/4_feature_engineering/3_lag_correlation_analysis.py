"""
HVAC Lag Correlation Analysis Script
Analyzes time-lagged relationships between variables
"""

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from pathlib import Path
import configparser
from datetime import datetime
import warnings
import sys

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None

warnings.filterwarnings('ignore')


class LagCorrelationAnalyzer:
    """Analyzes time-lagged correlations in HVAC data"""
    
    def __init__(self, config_path='config.ini'):
        """Initialize with config file path"""
        # Allow inline comments in config values
        self.config = configparser.ConfigParser(inline_comment_prefixes=(';', '#'))
        self.plot_records = []
        self.analysis_start = None
        self.analysis_end = None
        self.analysis_period_text = ""
        
        # Handle path resolution
        config_file = Path(config_path)
        
        if not config_file.is_absolute():
            if config_file.exists():
                config_file = config_file.resolve()
            else:
                script_dir = Path(__file__).parent if '__file__' in globals() else Path.cwd()
                parent_config = script_dir.parent / config_path
                
                if parent_config.exists():
                    config_file = parent_config
                elif not config_file.exists():
                    project_root = Path.cwd()
                    while project_root != project_root.parent:
                        test_config = project_root / config_path
                        if test_config.exists():
                            config_file = test_config
                            break
                        project_root = project_root.parent
        
        if not config_file.exists():
            raise FileNotFoundError(f"Config file not found: {config_path}")
        
        print(f"Reading config from: {config_file}")
        self.config.read(config_file, encoding='utf-8')
        self.config_dir = config_file.parent
        
        self.load_settings()
        
    def load_settings(self):
        """Load settings from config file"""
        # Global settings
        self.building_id = self.config.get('global', 'building_id')
        self.ahu_unit = self.config.get('global', 'ahu_unit')
        self.season = self.config.get('global', 'season')
        self.year = self.config.get('global', 'year')
        
        # Lag correlation settings
        lag_section = 'lag_correlation'
        self.enable_lag = self.config.getboolean(lag_section, 'enable_lag_correlation', fallback=True)
        
        # Parse lag pairs
        lag_pairs_str = self.config.get(lag_section, 'lag_pairs', 
                                        fallback='Temperatura Esterna:Temperatura Ripresa')
        self.lag_pairs = []
        for pair in lag_pairs_str.split(','):
            if ':' in pair:
                source, target = pair.strip().split(':')
                self.lag_pairs.append((source.strip(), target.strip()))
        
        # Lag settings
        self.max_lag_hours = self.config.getint(lag_section, 'max_lag_hours', fallback=6)
        self.time_resolution_minutes = self.config.getint(lag_section, 'time_resolution_minutes', fallback=30)
        
        # Calculate max lag in time steps
        self.max_lag_steps = int(self.max_lag_hours * 60 / self.time_resolution_minutes)
        
        # Correlation method and options
        self.corr_method = self.config.get(lag_section, 'method', fallback='pearson')
        self.twosided = self.config.getboolean(lag_section, 'two_sided', fallback=True)
        self.detrend_mode = self.config.get(lag_section, 'detrend', fallback='none')
        self.min_samples_per_lag = self.config.getint(lag_section, 'min_samples_per_lag', fallback=48)
        
        # Data filtering
        self.use_interpolated = self.config.getboolean(lag_section, 'use_interpolated', fallback=False)
        self.hour_start = self.config.getint(lag_section, 'hour_start', fallback=0)
        self.hour_end = self.config.getint(lag_section, 'hour_end', fallback=24)
        self.only_weekdays = self.config.getboolean(lag_section, 'only_weekdays', fallback=False)
        
        # Timezone
        self.timezone = self.config.get('global', 'timezone', fallback='Europe/Rome')
        
        # Visualization
        self.create_plots = self.config.getboolean(lag_section, 'create_lag_plots', fallback=True)
        self.mark_optimal = self.config.getboolean(lag_section, 'mark_optimal_lag', fallback=True)
        self.annotate_values = self.config.getboolean(lag_section, 'annotate_values', fallback=True)
        
        # Analysis options
        self.identify_response = self.config.getboolean(lag_section, 'identify_response_time', fallback=True)
        self.identify_latency = self.config.getboolean(lag_section, 'identify_control_latency', fallback=True)
        
        # Create base name FIRST (needed for paths)
        self.base_name = f"{self.building_id}_{self.ahu_unit}_{self.season}{self.year}"
        
        # Output settings - CSV files go to processed_data
        processed_folder = self._resolve_path(self.config.get('paths', 'processed_folder'))
        self.csv_output_dir = processed_folder / 'correlation_analysis' / self.base_name
        self.csv_output_dir.mkdir(parents=True, exist_ok=True)
        
        # Plots and reports go to plots folder
        plots_folder = self._resolve_path(self.config.get('paths', 'plots_folder', fallback='plots'))
        self.plots_output_dir = plots_folder / 'correlation_analysis' / self.base_name
        self.plots_output_dir.mkdir(parents=True, exist_ok=True)
        
        self.plot_format = self.config.get(lag_section, 'lag_plot_format', fallback='png')
        self.plot_dpi = self.config.getint(lag_section, 'lag_plot_dpi', fallback=300)
        
        # File paths - Use selected features for focused correlation analysis
        
        # 1. Find selected features CSV (output from feature importance analyzer)
        selected_features_filename = f"{self.base_name}_selected_features.csv"
        selected_candidates = [
            processed_folder / 'feature_analysis' / self.base_name / selected_features_filename,
            processed_folder / 'feature_engineering' / self.base_name / selected_features_filename,
            processed_folder / 'feature_importance' / self.base_name / selected_features_filename,
        ]
        
        self.selected_features_file = None
        for candidate in selected_candidates:
            if candidate.exists():
                self.selected_features_file = candidate
                print(f"  ✓ Found selected features: {candidate}")
                break
                self.selected_features_file = candidate
                break
        
        # 2. Find full features data file
        features_filename = f"{self.base_name}_features.csv"
        feature_candidates = [
            processed_folder / 'feature_engineering' / self.base_name / features_filename,
            processed_folder / 'interpolation' / self.base_name / features_filename,
            processed_folder / 'interpolation' / features_filename,
        ]
        
        # Fallback to interpolated if features not found
        interp_filename = f"{self.base_name}_interpolated.csv"
        interp_candidates = [
            processed_folder / 'interpolation' / self.base_name / interp_filename,
            processed_folder / 'interpolation' / interp_filename,
        ]
        
        all_candidates = feature_candidates + interp_candidates
        
        self.data_file = all_candidates[0]
        for candidate in all_candidates:
            if candidate.exists():
                self.data_file = candidate
                break
        
    def _resolve_path(self, path_str):
        """Resolve path relative to config directory"""
        path = Path(path_str)
        if not path.is_absolute():
            path = self.config_dir / path
        return path
    
    def _format_date(self, value):
        """Format datetime/date values for display"""
        if value is None:
            return "N/D"
        if isinstance(value, pd.Timestamp):
            value = value.to_pydatetime()
        if hasattr(value, 'strftime'):
            return value.strftime('%d/%m/%Y')
        return str(value)
    
    def _update_analysis_period(self, df):
        """Update analysis period text"""
        if df is None or df.empty:
            self.analysis_period_text = "Periodo di analisi: Dati non disponibili"
            return
        
        self.analysis_start = df.index.min()
        self.analysis_end = df.index.max()
        self.analysis_period_text = f"Periodo di analisi: {self._format_date(self.analysis_start)} - {self._format_date(self.analysis_end)}"
    
    def _prep_series(self, series):
        """
        Preprocess series for correlation with optional detrending.
        
        Args:
            series: pandas Series to preprocess
            
        Returns:
            Preprocessed series with NaN removed
        """
        s = series.copy().dropna()
        
        if self.detrend_mode == 'demean':
            s = s - s.mean()
        elif self.detrend_mode == 'diff':
            s = s.diff().dropna()
        # else: 'none', return as-is
        
        return s
    
    def _smart_resample(self, df, freq):
        """
        Resample with smart aggregation: mean for numeric, ffill for binary/control.
        
        Args:
            df: DataFrame to resample
            freq: Target frequency string (e.g., '30min')
            
        Returns:
            Resampled DataFrame
        """
        # Identify column types for smart aggregation
        binary_cols = []
        numeric_cols = []
        
        for col in df.columns:
            if df[col].dtype in ['bool', 'boolean']:
                binary_cols.append(col)
            elif col.lower().startswith('is_') or col.lower().endswith('_flag'):
                # Heuristic: columns starting with 'is_' or ending with '_flag' are binary
                binary_cols.append(col)
            elif 'setpoint' in col.lower() or 'mode' in col.lower():
                # Setpoints and modes should use forward fill
                binary_cols.append(col)
            elif pd.api.types.is_numeric_dtype(df[col]):
                numeric_cols.append(col)
            else:
                # Non-numeric, non-binary → skip or forward fill
                binary_cols.append(col)
        
        # Resample with appropriate aggregation
        resampled_parts = []
        
        if numeric_cols:
            resampled_parts.append(df[numeric_cols].resample(freq).mean())
        
        if binary_cols:
            resampled_parts.append(df[binary_cols].resample(freq).ffill())
        
        # Combine
        if resampled_parts:
            result = pd.concat(resampled_parts, axis=1)
            # Restore original column order
            result = result[df.columns]
            return result
        
        return df
    
    def load_data(self):
        """Load data with timezone validation, frequency enforcement, and filtering"""
        print(f"Loading data from: {self.data_file}")
        
        if not self.data_file.exists():
            raise FileNotFoundError(f"Data file not found: {self.data_file}")
        
        df = pd.read_csv(self.data_file)
        
        # Parse time column
        time_col = self.config.get('data', 'time_column', fallback='Time')
        if time_col not in df.columns:
            raise ValueError(f"Time column '{time_col}' not found")
        
        # Parse datetime with UTC handling for timezone-aware strings
        df[time_col] = pd.to_datetime(df[time_col], utc=True)
        df.set_index(time_col, inplace=True)
        
        print(f"Loaded {len(df)} records from {df.index.min()} to {df.index.max()}")
        
        # Timezone handling (Critical for DST and time drift)
        if df.index.tz is None:
            df.index = df.index.tz_localize(self.timezone, ambiguous='NaT', nonexistent='shift_forward')
            # Drop any NaT from ambiguous times
            df = df[~df.index.isna()]
            print(f"  ✓ Localized to {self.timezone}")
        else:
            df.index = df.index.tz_convert(self.timezone)
            print(f"  ✓ Converted to {self.timezone}")
        
        # Frequency validation and enforcement
        try:
            detected_freq = pd.infer_freq(df.index[:min(100, len(df))])
        except Exception:
            detected_freq = None
        
        expected_freq = f'{self.time_resolution_minutes}min'
        expected_td = pd.Timedelta(minutes=self.time_resolution_minutes)
        
        if detected_freq:
            detected_td = pd.Timedelta(detected_freq)
            print(f"  ℹ Detected frequency: {detected_freq}")
            
            if detected_td != expected_td:
                print(f"  ⚠ Frequency mismatch: expected {expected_freq}, got {detected_freq}")
                print(f"  → Resampling to {expected_freq} with smart aggregation")
                df = self._smart_resample(df, expected_freq)
        else:
            print(f"  ⚠ Could not infer frequency, resampling to {expected_freq}")
            df = self._smart_resample(df, expected_freq)
        
        # Exclude interpolated points (avoid overstating correlations on fabricated data)
        if not self.use_interpolated:
            mask_cols = [c for c in df.columns if c.endswith('_is_interpolated')]
            if mask_cols:
                not_interp = ~(df[mask_cols].any(axis=1))
                n_before = len(df)
                df = df.loc[not_interp]
                n_excluded = n_before - len(df)
                print(f"  ✓ Excluded {n_excluded} interpolated points ({n_excluded/n_before*100:.1f}%)")
                # Drop interpolation mask columns
                df = df.drop(columns=mask_cols)
        
        # Apply time-of-day filters (occupancy windows)
        if self.hour_start > 0 or self.hour_end < 24:
            time_mask = (df.index.hour >= self.hour_start) & (df.index.hour < self.hour_end)
            n_before = len(df)
            df = df[time_mask]
            print(f"  ✓ Filtered to hours {self.hour_start:02d}:00-{self.hour_end:02d}:00 ({len(df)}/{n_before} records)")
        
        # Apply weekday filter
        if self.only_weekdays:
            weekday_mask = df.index.dayofweek < 5
            n_before = len(df)
            df = df[weekday_mask]
            print(f"  ✓ Filtered to weekdays only ({len(df)}/{n_before} records)")
        
        # Filter to selected features (if available)
        if self.selected_features_file and self.selected_features_file.exists():
            try:
                selected_df = pd.read_csv(self.selected_features_file)
                selected_features = selected_df['feature'].tolist()
                
                # Keep only columns that exist in both selected features and current dataframe
                available_features = [f for f in selected_features if f in df.columns]
                missing_features = [f for f in selected_features if f not in df.columns]
                
                if available_features:
                    df = df[available_features]
                    print(f"  ✓ Filtered to {len(available_features)} selected features from importance analysis")
                    if missing_features:
                        print(f"    ℹ Skipped {len(missing_features)} missing features")
                else:
                    print(f"  ⚠ Warning: No selected features found in data, using all columns")
            except Exception as e:
                print(f"  ⚠ Warning: Could not load selected features ({e}), using all columns")
        else:
            print(f"  ℹ No selected features file found, using all available columns")
        
        print(f"Final dataset: {len(df)} records × {len(df.columns)} features from {df.index.min()} to {df.index.max()}")
        self._update_analysis_period(df)
        return df
    
    def compute_lag_correlation(self, df, source_col, target_col, max_lag):
        """
        Compute correlation at different lags with optional detrending and two-sided scanning.
        
        Args:
            df: DataFrame with time series data
            source_col: Column name for lagged variable
            target_col: Column name for reference variable
            max_lag: Maximum lag in steps (will scan -max_lag to +max_lag if two_sided=True)
            
        Returns:
            lags_hours: List of lag values in hours
            correlations: List of correlation coefficients
            n_samples: List of effective sample sizes per lag
        """
        correlations = []
        lags_hours = []
        n_samples = []
        
        # Check if columns exist
        if source_col not in df.columns:
            print(f"Warning: Source column '{source_col}' not found")
            return None, None, None
        if target_col not in df.columns:
            print(f"Warning: Target column '{target_col}' not found")
            return None, None, None
        
        # Preprocess series (detrending)
        source = self._prep_series(df[source_col])
        target = self._prep_series(df[target_col])
        
        # Re-align after preprocessing
        common_idx = source.index.intersection(target.index)
        source = source.loc[common_idx]
        target = target.loc[common_idx]
        
        # Determine lag range
        if self.twosided:
            lag_range = range(-max_lag, max_lag + 1)
        else:
            lag_range = range(0, max_lag + 1)
        
        # Compute correlations for each lag
        for lag in lag_range:
            # Positive lag: source leads target by lag steps
            # Negative lag: target leads source by |lag| steps
            source_shifted = source.shift(lag)
            
            # Calculate correlation with target
            valid_data = pd.DataFrame({
                'source': source_shifted,
                'target': target
            }).dropna()
            
            n_eff = len(valid_data)
            lag_hours = lag * self.time_resolution_minutes / 60
            
            # Sample size validation
            if n_eff >= self.min_samples_per_lag:
                if self.corr_method == 'spearman':
                    corr = valid_data['source'].corr(valid_data['target'], method='spearman')
                else:
                    corr = valid_data['source'].corr(valid_data['target'], method='pearson')
                
                correlations.append(corr)
                lags_hours.append(lag_hours)
                n_samples.append(n_eff)
            else:
                # Skip lags with insufficient data
                correlations.append(np.nan)
                lags_hours.append(lag_hours)
                n_samples.append(n_eff)
        
        return lags_hours, correlations, n_samples
    
    def find_optimal_lag(self, lags, correlations):
        """Find lag with highest absolute correlation"""
        correlations_arr = np.array(correlations)
        valid_mask = ~np.isnan(correlations_arr)
        
        if not valid_mask.any():
            return None, None
        
        abs_corr = np.abs(correlations_arr[valid_mask])
        max_idx = np.argmax(abs_corr)
        
        # Get indices of valid correlations
        valid_indices = np.where(valid_mask)[0]
        optimal_idx = valid_indices[max_idx]
        
        optimal_lag = lags[optimal_idx]
        optimal_corr = correlations[optimal_idx]
        
        return optimal_lag, optimal_corr
    
    def plot_lag_analysis(self, results):
        """Create lag correlation plots"""
        n_pairs = len(results)
        
        if n_pairs == 0:
            print("No valid lag correlations to plot")
            return
        
        # Create subplot layout
        n_cols = 2
        n_rows = (n_pairs + 1) // 2
        
        fig, axes = plt.subplots(n_rows, n_cols, figsize=(14, 4 * n_rows))
        if n_pairs == 1:
            axes = np.array([axes])
        axes = axes.flatten()
        
        for idx, (pair_name, data) in enumerate(results.items()):
            ax = axes[idx]
            
            lags = data['lags']
            correlations = data['correlations']
            n_samples = data.get('n_samples', [])
            optimal_lag = data['optimal_lag']
            optimal_corr = data['optimal_corr']
            
            # Convert to arrays for safe processing
            lags_arr = np.array(lags)
            corrs_arr = np.array(correlations)
            
            # Plot correlation vs lag
            ax.plot(lags, correlations, 'o-', linewidth=2, markersize=6, label='Correlation')
            ax.axhline(y=0, color='gray', linestyle='--', alpha=0.5)
            
            # Add vertical line at zero lag for two-sided plots
            if self.twosided:
                ax.axvline(x=0, color='black', linestyle='-', linewidth=0.8, alpha=0.3)
            
            ax.grid(True, alpha=0.3)
            
            # Mark optimal lag
            if self.mark_optimal and optimal_lag is not None:
                ax.axvline(x=optimal_lag, color='red', linestyle='--', linewidth=2, 
                          label=f'Optimal lag: {optimal_lag:.1f}h')
                ax.plot(optimal_lag, optimal_corr, 'r*', markersize=15)
                
                if self.annotate_values:
                    ax.annotate(f'r={optimal_corr:.3f}', 
                              xy=(optimal_lag, optimal_corr),
                              xytext=(10, 10), textcoords='offset points',
                              fontsize=10, fontweight='bold',
                              bbox=dict(boxstyle='round', facecolor='yellow', alpha=0.7))
            
            # Axis labels with semantic meaning for two-sided
            if self.twosided:
                xlabel = 'Ritardo Temporale (ore)\n← Target leads | Source leads →'
            else:
                xlabel = 'Ritardo Temporale (ore)'
            
            ax.set_xlabel(xlabel, fontsize=11)
            ax.set_ylabel('Coefficiente di Correlazione', fontsize=11)
            ax.set_title(pair_name, fontsize=12, fontweight='bold', color='#003366')
            ax.legend(fontsize=9)
            
            # Robust y-limits based on finite values
            finite_corrs = corrs_arr[np.isfinite(corrs_arr)]
            if len(finite_corrs) > 0:
                ymin = max(-1.0, np.min(finite_corrs) - 0.1)
                ymax = min(1.0, np.max(finite_corrs) + 0.1)
                ax.set_ylim(ymin, ymax)
            else:
                ax.set_ylim(-1, 1)
            
            # Warn about sample size if any lags have insufficient data
            if n_samples:
                n_arr = np.array(n_samples)
                n_insufficient = np.sum(n_arr < self.min_samples_per_lag)
                if n_insufficient > 0:
                    ax.text(0.02, 0.98, f'⚠ {n_insufficient} lags with n < {self.min_samples_per_lag}',
                           transform=ax.transAxes, fontsize=8, va='top', 
                           bbox=dict(boxstyle='round', facecolor='wheat', alpha=0.5))
        
        # Hide unused subplots
        for idx in range(n_pairs, len(axes)):
            axes[idx].set_visible(False)
        
        fig.suptitle(f'Analisi di Correlazione con Ritardo Temporale\n{self.building_id} - {self.ahu_unit} - {self.season} {self.year}',
                    fontsize=14, fontweight='bold', y=0.98)
        
        # Add date subtitle
        if self.analysis_period_text:
            fig.text(0.5, 0.96, self.analysis_period_text, ha='center', 
                    fontsize=10, style='italic', color='gray')
        
        plt.tight_layout(rect=[0, 0, 1, 0.94])
        
        # Save plot to plots folder (use pre-created directory)
        filename = f"lag_correlation_{self.building_id}_{self.ahu_unit}_{self.season}{self.year}.{self.plot_format}"
        filepath = self.plots_output_dir / filename
        plt.savefig(filepath, dpi=self.plot_dpi, bbox_inches='tight')
        print(f"Saved lag correlation plot: {filepath}")
        
        self.plot_records.append({
            'title': 'Analisi di Correlazione con Ritardo Temporale',
            'path': filepath,
            'description': 'Analizza come le variabili si influenzano nel tempo, identificando ritardi ottimali per il controllo predittivo. I grafici mostrano la correlazione tra coppie di variabili a diversi ritardi temporali, evidenziando il ritardo ottimale (linea rossa tratteggiata) dove la correlazione è massima. Questo aiuta a comprendere i tempi di risposta termica del sistema e la latenza del sistema di controllo.'
        })
        
        plt.close()
    
    def export_csv_results(self, results):
        """Export lag correlation results to CSV files"""
        # Use pre-created CSV output directory (processed_data/correlation_analysis/base_name/)
        csv_dir = self.csv_output_dir
        
        # Export per-pair detailed results
        for idx, (pair_name, data) in enumerate(results.items(), 1):
            # Use shorter filenames to avoid Windows path length issues
            safe_source = data['source'].replace(' ', '_')[:30]  # Truncate long names
            safe_target = data['target'].replace(' ', '_')[:30]
            filename = f"lag_corr_{idx}_{safe_source}_to_{safe_target}.csv"
            filepath = csv_dir / filename
            
            # Create DataFrame with lag, correlation, sample size
            df_pair = pd.DataFrame({
                'lag_hours': data['lags'],
                'correlation': data['correlations'],
                'n_samples': data.get('n_samples', [None] * len(data['lags']))
            })
            
            # Add metadata as comment at top
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(f"# Lag Correlation: {pair_name}\n")
                f.write(f"# Source: {data['source']}\n")
                f.write(f"# Target: {data['target']}\n")
            df_pair.to_csv(filepath, index=False, float_format='%.6f', mode='a', encoding='utf-8')
            print(f"  ✓ Saved {filepath.name} ({pair_name})")
        
        # Export summary with optimal lags
        summary_data = []
        for pair_name, data in results.items():
            # Find n_samples at optimal lag (not just first lag)
            optimal_lag = data.get('optimal_lag')
            n_samples_list = data.get('n_samples', [])
            lags_list = data.get('lags', [])
            
            n_at_optimal = None
            if optimal_lag is not None and lags_list and n_samples_list:
                try:
                    # Find index of optimal lag
                    optimal_idx = lags_list.index(optimal_lag)
                    n_at_optimal = n_samples_list[optimal_idx]
                except (ValueError, IndexError):
                    n_at_optimal = None
            
            summary_data.append({
                'pair': pair_name,
                'source': data['source'],
                'target': data['target'],
                'optimal_lag_hours': optimal_lag,
                'optimal_correlation': data.get('optimal_corr'),
                'n_effective_samples': n_at_optimal
            })
        
        df_summary = pd.DataFrame(summary_data)
        summary_filename = f"lag_correlation_summary_{self.building_id}_{self.ahu_unit}_{self.season}{self.year}.csv"
        summary_filepath = csv_dir / summary_filename
        df_summary.to_csv(summary_filepath, index=False, float_format='%.6f')
        print(f"  ✓ Saved summary: {summary_filepath.name}")
    
    def generate_findings_report(self, results):
        """Generate text report of findings"""
        print("\n" + "="*70)
        print("LAG CORRELATION FINDINGS")
        print("="*70)
        
        findings = []
        
        for pair_name, data in results.items():
            optimal_lag = data['optimal_lag']
            optimal_corr = data['optimal_corr']
            
            if optimal_lag is None:
                print(f"\n{pair_name}:")
                print("  No valid correlation found")
                continue
            
            print(f"\n{pair_name}:")
            print(f"  Optimal lag: {optimal_lag:.1f} hours ({int(optimal_lag * 60 / self.time_resolution_minutes)} time steps)")
            print(f"  Maximum correlation: {optimal_corr:.3f}")
            
            # Interpret the finding
            if 'Esterna' in pair_name and 'Ripresa' in pair_name:
                print(f"  → Thermal response time: ~{optimal_lag:.1f} hours")
                findings.append({
                    'type': 'Thermal Response',
                    'lag': optimal_lag,
                    'correlation': optimal_corr,
                    'interpretation': f'Outdoor conditions affect indoor temperature after {optimal_lag:.1f} hours'
                })
            
            elif 'error' in pair_name.lower() and 'Modul' in pair_name:
                print(f"  → Control system latency: ~{optimal_lag:.1f} hours")
                findings.append({
                    'type': 'Control Latency',
                    'lag': optimal_lag,
                    'correlation': optimal_corr,
                    'interpretation': f'Control response to setpoint error takes {optimal_lag:.1f} hours'
                })
            
            elif 'Mandata' in pair_name and 'Ripresa' in pair_name:
                print(f"  → Air distribution time: ~{optimal_lag:.1f} hours")
                findings.append({
                    'type': 'Air Distribution',
                    'lag': optimal_lag,
                    'correlation': optimal_corr,
                    'interpretation': f'Supply air affects return air after {optimal_lag:.1f} hours'
                })
        
        # Save findings to CSV (in processed_data folder)
        if findings:
            findings_df = pd.DataFrame(findings)
            findings_file = self.csv_output_dir / f"lag_findings_{self.base_name}.csv"
            findings_df.to_csv(findings_file, index=False)
            print(f"\nSaved findings: {findings_file}")
        
        return findings
    
    def create_report_document(self, results):
        """Create Word document with plots and analysis results"""
        if Document is None:
            print("python-docx not available: skipping Word document")
            return
        
        doc = Document()
        
        # Main Title (centered, bold)
        title = doc.add_heading('Analisi di Correlazione HVAC', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Metadata section
        doc.add_paragraph(f"Edificio: {self.building_id}", style='Heading 2')
        doc.add_paragraph(f"Unità UTA: {self.ahu_unit}")
        doc.add_paragraph(f"Stagione: {self.season} {self.year}")
        doc.add_paragraph(f"{self.analysis_period_text}")
        doc.add_paragraph(f"Report Generato: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph("")
        
        # Methods section
        doc.add_heading('Metodi di Analisi', level=2)
        doc.add_paragraph(f"Metodo di Correlazione: {self.corr_method.capitalize()}")
        
        lag_type = "bidirezionale (-{} a +{} ore)".format(
            self.max_lag_hours, self.max_lag_hours
        ) if self.twosided else f"unidirezionale (0 a {self.max_lag_hours} ore)"
        doc.add_paragraph(f"Tipo di Scansione Lag: {lag_type}")
        
        if self.detrend_mode != 'none':
            doc.add_paragraph(f"Preprocessing: Detrending '{self.detrend_mode}'")
        
        doc.add_paragraph(f"Campioni Minimi per Lag: {self.min_samples_per_lag}")
        
        if not self.use_interpolated:
            doc.add_paragraph("Punti interpolati esclusi dall'analisi")
        
        if self.hour_start > 0 or self.hour_end < 24:
            doc.add_paragraph(f"Filtro Orario: {self.hour_start:02d}:00 - {self.hour_end:02d}:00")
        
        if self.only_weekdays:
            doc.add_paragraph("Solo giorni feriali (Lun-Ven)")
        
        doc.add_paragraph("")
        
        # Add each plot with structure requested
        for plot_info in self.plot_records:
            # Row 1: Plot title (centered, bold, dark blue)
            heading = doc.add_heading(plot_info['title'], level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0, 51, 102)  # Dark blue
            
            # Row 2: Plot image (centered)
            if plot_info['path'].exists():
                doc.add_paragraph()  # Spacing
                doc.add_picture(str(plot_info['path']), width=Inches(6.5))
                last_paragraph = doc.paragraphs[-1]
                last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Row 3: Description section
            doc.add_paragraph()  # Spacing
            desc_heading = doc.add_paragraph()
            desc_run = desc_heading.add_run('Descrizione:')
            desc_run.bold = True
            desc_run.font.size = Pt(11)
            
            desc_paragraph = doc.add_paragraph(plot_info['description'])
            desc_paragraph.style = 'Normal'
            
            # Add detailed results for lag correlation
            if 'lag' in plot_info['title'].lower():
                doc.add_paragraph()
                results_heading = doc.add_paragraph()
                results_run = results_heading.add_run('Risultati Dettagliati:')
                results_run.bold = True
                results_run.font.size = Pt(11)
                
                for pair_name, data in results.items():
                    optimal_lag = data['optimal_lag']
                    optimal_corr = data['optimal_corr']
                    
                    pair_para = doc.add_paragraph()
                    pair_para.add_run(f"• {pair_name}: ").bold = True
                    
                    if optimal_lag is not None:
                        pair_para.add_run(f"Ritardo ottimale = {optimal_lag:.1f} ore, Correlazione = {optimal_corr:.3f}")
                        
                        # Interpretation
                        if abs(optimal_corr) > 0.7:
                            strength = "Forte"
                        elif abs(optimal_corr) > 0.4:
                            strength = "Moderata"
                        else:
                            strength = "Debole"
                        
                        direction = "positiva" if optimal_corr > 0 else "negativa"
                        pair_para.add_run(f" ({strength} correlazione {direction})")
                    else:
                        pair_para.add_run("Nessuna correlazione significativa trovata")
            
            doc.add_page_break()
        
        # Recommendations section
        doc.add_heading('Raccomandazioni per il Controllo Predittivo', level=1)
        doc.add_paragraph("Sulla base dell'analisi dei ritardi temporali, si consiglia di:")
        doc.add_paragraph("• Utilizzare la temperatura esterna con il ritardo identificato per il controllo feed-forward", style='List Bullet')
        doc.add_paragraph("• Regolare il tempo di risposta del sistema di controllo se la latenza è eccessiva", style='List Bullet')
        doc.add_paragraph("• Implementare il Controllo Predittivo del Modello (MPC) con orizzonti di predizione ottimali", style='List Bullet')
        doc.add_paragraph("• Considerare l'inerzia termica dell'edificio nella progettazione del controllo", style='List Bullet')
        
        # Save document to plots folder (use pre-created directory)
        doc_file = self.plots_output_dir / f"correlation_analysis_report_{self.base_name}.docx"
        doc.save(str(doc_file))
        print(f"Saved report document: {doc_file}")
    
    def run_analysis(self):
        """Run complete lag correlation analysis"""
        if not self.enable_lag:
            print("Lag correlation analysis is disabled in config")
            return
        
        print("\n" + "="*70)
        print("LAG CORRELATION ANALYSIS")
        print("="*70)
        print(f"Building: {self.building_id}")
        print(f"AHU Unit: {self.ahu_unit}")
        print(f"Season: {self.season} {self.year}")
        print(f"Max lag: {self.max_lag_hours} hours ({self.max_lag_steps} steps)")
        print("="*70)
        
        # Load data
        try:
            df = self.load_data()
        except Exception as e:
            print(f"\nFailed to load data: {e}")
            return
        
        # Compute lag correlations for each pair
        results = {}
        
        print(f"\nComputing lag correlations for {len(self.lag_pairs)} pairs...")
        
        for source_col, target_col in self.lag_pairs:
            pair_name = f"{source_col} → {target_col}"
            print(f"\n  {pair_name}")
            
            lags, correlations, n_samples = self.compute_lag_correlation(
                df, source_col, target_col, self.max_lag_steps
            )
            
            if lags is not None:
                optimal_lag, optimal_corr = self.find_optimal_lag(lags, correlations)
                
                results[pair_name] = {
                    'lags': lags,
                    'correlations': correlations,
                    'n_samples': n_samples,
                    'optimal_lag': optimal_lag,
                    'optimal_corr': optimal_corr,
                    'source': source_col,
                    'target': target_col
                }
                
                if optimal_lag is not None:
                    print(f"    Optimal lag: {optimal_lag:.1f}h (r={optimal_corr:.3f})")
        
        # Export CSV results
        if results:
            print("\nExporting CSV results...")
            self.export_csv_results(results)
        
        # Create visualizations
        if self.create_plots and results:
            print("\nCreating lag correlation plots...")
            self.plot_lag_analysis(results)
        
        # Generate findings report
        if results:
            findings = self.generate_findings_report(results)
            self.create_report_document(results)
        
        print("\n" + "="*70)
        print("LAG CORRELATION ANALYSIS COMPLETED")
        print("="*70)


if __name__ == "__main__":
    config_path = sys.argv[1] if len(sys.argv) > 1 else 'config.ini'
    
    print(f"HVAC Lag Correlation Analysis")
    print(f"=" * 70)
    print(f"Python version: {sys.version}")
    print(f"Working directory: {Path.cwd()}")
    print(f"=" * 70)
    
    try:
        analyzer = LagCorrelationAnalyzer(config_path)
        analyzer.run_analysis()
    except Exception as e:
        print(f"\n✗ Error: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)